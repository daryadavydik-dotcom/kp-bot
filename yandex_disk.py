import csv
from io import BytesIO, StringIO

import httpx
import openpyxl


class YandexDiskClient:
    _BASE = "https://cloud-api.yandex.net/v1/disk"

    def __init__(self, token: str):
        self._headers = {"Authorization": f"OAuth {token}"}

    async def download_file(self, path: str) -> bytes:
        async with httpx.AsyncClient(timeout=30) as client:
            resp = await client.get(
                f"{self._BASE}/resources/download",
                headers=self._headers,
                params={"path": path},
            )
            resp.raise_for_status()
            file_resp = await client.get(resp.json()["href"], follow_redirects=True)
            file_resp.raise_for_status()
            return file_resp.content

    async def upload_file(self, path: str, content: bytes) -> None:
        async with httpx.AsyncClient(timeout=60) as client:
            resp = await client.get(
                f"{self._BASE}/resources/upload",
                headers=self._headers,
                params={"path": path, "overwrite": "true"},
            )
            resp.raise_for_status()
            put_resp = await client.put(resp.json()["href"], content=content)
            put_resp.raise_for_status()

    async def publish_and_get_url(self, path: str) -> str:
        async with httpx.AsyncClient(timeout=30) as client:
            pub_resp = await client.put(
                f"{self._BASE}/resources/publish",
                headers=self._headers,
                params={"path": path},
            )
            pub_resp.raise_for_status()

            info_resp = await client.get(
                f"{self._BASE}/resources",
                headers=self._headers,
                params={"path": path, "fields": "public_url"},
            )
            info_resp.raise_for_status()
            return info_resp.json().get("public_url", "")

    async def create_folder(self, path: str) -> None:
        async with httpx.AsyncClient(timeout=15) as client:
            resp = await client.put(
                f"{self._BASE}/resources",
                headers=self._headers,
                params={"path": path},
            )
            if resp.status_code not in (201, 409):  # 409 = already exists
                resp.raise_for_status()

    async def load_nomenclature(self, path: str) -> list[dict]:
        content = await self.download_file(path)
        if path.lower().endswith(".csv"):
            return self._parse_csv(content)
        return self._parse_excel(content)

    @staticmethod
    def _parse_excel(content: bytes) -> list[dict]:
        wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
        ws = wb.active
        headers = [str(cell.value) if cell.value is not None else "" for cell in ws[1]]
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            row_dict = {h: v for h, v in zip(headers, row) if h}
            if any(v is not None for v in row_dict.values()):
                rows.append(row_dict)
        return rows

    @staticmethod
    def _parse_csv(content: bytes) -> list[dict]:
        text = content.decode("utf-8-sig")
        reader = csv.DictReader(StringIO(text))
        return list(reader)
